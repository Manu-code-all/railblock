"""
SIH 2026 — PS 27 Automatic Block Planning (Indian Railways)
Fills the official SIH template: concise text on the left, native
PowerPoint diagrams on the right (vector, fully editable).

Run:  python fill_ppt.py
Then: open the .pptx in PowerPoint -> File -> Export -> Create PDF
"""
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

import os
HERE = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(HERE, "SIH2026-IDEA-Presentation-Format.pptx")
DEST = os.path.join(HERE, "SIH2026_PS27_RailBlock.pptx")

# ── Team details ────────────────────────────────────────────
TEAM_NAME = "Anomaly"
TEAM_ID   = ""                        # fill once the portal assigns it
INSTITUTE = "Galgotias University, Uttar Pradesh"
# ────────────────────────────────────────────────────────────

# ── Palette ─────────────────────────────────────────────────
NAVY   = "1F3864"
BLUE   = "2E6DA4"
STEEL  = "4A7BA7"
GREEN  = "1E8449"
GREENL = "E8F5EC"
RED    = "C0392B"
REDL   = "FDEDEC"
AMBER  = "CA8A04"
AMBERL = "FEF6E0"
GREY   = "3A3A3A"
GREYL  = "EEF1F5"
WHITE  = "FFFFFF"

# Department colours (used in the Gantt)
C_ENG  = "2E6DA4"
C_SIG  = "1E8449"
C_TRD  = "CA8A04"

prs = Presentation(SRC)


# ══════════════════════════════════════════════════════════════
#  Text helpers
# ══════════════════════════════════════════════════════════════
def _clean_para(p, space_before=0):
    """Strip inherited bullets/indent from a paragraph."""
    pPr = p._p.get_or_add_pPr()
    if space_before:
        p.space_before = Pt(space_before)
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def para(tf, text, size=11.5, bold=False, color=GREY, space_before=0,
         first=False, italic=False, align=PP_ALIGN.LEFT):
    if first:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    _clean_para(p, space_before)
    p.alignment = align
    if text:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def body(shape, blocks, x, y, w, h):
    """Fill a template textbox with a compact head/bullet block list."""
    shape.left, shape.top = Inches(x), Inches(y)
    shape.width, shape.height = Inches(w), Inches(h)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    first = True
    for kind, text in blocks:
        if kind == "head":
            para(tf, text.upper(), size=13, bold=True, color=NAVY,
                 space_before=0 if first else 15, first=first)
        elif kind == "b":
            para(tf, "•  " + text, size=12, color=GREY,
                 space_before=5, first=first)
        elif kind == "bg":
            para(tf, "•  " + text, size=12, bold=True, color=GREEN,
                 space_before=5, first=first)
        elif kind == "note":
            para(tf, text, size=10.5, italic=True, color=STEEL,
                 space_before=8, first=first)
        first = False


def set_text(shape, text, size=None):
    tf = shape.text_frame
    tf.word_wrap = False
    tf.clear()
    p = tf.paragraphs[0]
    r = p.runs[0] if p.runs else p.add_run()
    r.text = text
    if size:
        r.font.size = Pt(size)


def find(slide, prefix):
    for s in slide.shapes:
        if s.name.startswith(prefix):
            return s
    return None


# ══════════════════════════════════════════════════════════════
#  Diagram helpers  (native PowerPoint shapes)
# ══════════════════════════════════════════════════════════════
def box(slide, x, y, w, h, text="", fill=WHITE, line=NAVY, font=NAVY,
        size=10, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        lw=1.0, sub=None, sub_size=8.5, sub_color=None,
        align=PP_ALIGN.CENTER, pad=0.05):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False

    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.clear()
    para(tf, text, size=size, bold=bold, color=font,
         align=align, first=True)
    if sub:
        para(tf, sub, size=sub_size, color=sub_color or font, align=align)
    return sp


def arrow(slide, x, y, w, h, direction="down", fill=STEEL):
    shp = MSO_SHAPE.DOWN_ARROW if direction == "down" else MSO_SHAPE.RIGHT_ARROW
    sp = slide.shapes.add_shape(shp, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def label(slide, x, y, w, h, text, size=9, bold=False, color=NAVY,
          align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.clear()
    para(tf, text, size=size, bold=bold, color=color, align=align,
         italic=italic, first=True)
    return tb


def bar(slide, x, y, w, h, fill):
    """Flat Gantt bar."""
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


# ══════════════════════════════════════════════════════════════
#  Geometry
# ══════════════════════════════════════════════════════════════
LX, LW = 0.42, 5.55           # left text column
RX, RW = 6.35, 6.55           # right diagram column
TOP    = 1.12                 # content top


# ══════════════════════════════════════════════════════════════
#  SLIDE 1 — Title
# ══════════════════════════════════════════════════════════════
s1 = prs.slides[0]
tid = TEAM_ID.strip() or "(assigned on SIH portal)"

tb = find(s1, "TextBox 9")
tf = tb.text_frame
tf.clear()
tf.word_wrap = True
para(tf, "Problem Statement ID", size=10, color=STEEL, first=True)
para(tf, "SIH2026027", size=17, bold=True, color=NAVY)
para(tf, "Problem Statement Title", size=10, color=STEEL, space_before=10)
para(tf, "Automatic Block Planning for Indian Railways",
     size=14, bold=True, color=NAVY)
para(tf, "Theme  •  Smart Automation", size=11, color=GREY, space_before=10)
para(tf, "Category  •  Software", size=11, color=GREY, space_before=2)
para(tf, f"Team ID  •  {tid}", size=11, color=GREY, space_before=2)
para(tf, f"Team Name  •  {TEAM_NAME}", size=12, bold=True, color=NAVY,
     space_before=8)
para(tf, INSTITUTE, size=10, color=STEEL, space_before=2)


# ══════════════════════════════════════════════════════════════
#  SLIDE 2 — Proposed Solution
# ══════════════════════════════════════════════════════════════
s2 = prs.slides[1]
set_text(find(s2, "Title"), "PROPOSED SOLUTION")
set_text(find(s2, "Oval"), TEAM_NAME, size=12)

body(find(s2, "TextBox 8"), [
    ("head", "The problem"),
    ("b", "Engineering, Signalling and Traction each book track blocks in isolation"),
    ("b", "Blocks clash → windows wasted, maintenance deferred, trains delayed"),
    ("head", "Our solution — RailBlock"),
    ("b", "One engine ingests all three departments' demands plus the timetable"),
    ("b", "Returns a conflict-free block plan, weekly and monthly"),
    ("head", "Why it is different"),
    ("bg", "Proven, not predicted — CP-SAT certifies the plan is optimal"),
    ("b", "When demands cannot all fit, it names the exact conflicting requests"),
    ("b", "Priority slider re-solves the whole corridor in under 2 seconds"),
    ("b", "No training data, no model drift, every decision traceable"),
], LX, TOP, LW, 5.1)

# ── Diagram: demands -> solver -> two outcomes ──
label(s2, RX, TOP - 0.02, RW, 0.22,
      "MAINTENANCE DEMANDS", size=8.5, bold=True, color=STEEL,
      align=PP_ALIGN.CENTER)

dept = [("ENGINEERING", "track · civil", C_ENG),
        ("SIGNALLING", "interlocking", C_SIG),
        ("TRACTION", "OHE · power", C_TRD)]
bw, gap = 2.02, 0.24
for i, (nm, sb, col) in enumerate(dept):
    box(s2, RX + i * (bw + gap), TOP + 0.24, bw, 0.62, nm,
        fill=WHITE, line=col, font=col, size=9.5, bold=True,
        sub=sb, sub_size=7.5, sub_color=STEEL, lw=1.5)

for i in range(3):
    arrow(s2, RX + i * (bw + gap) + bw / 2 - 0.11, TOP + 0.92, 0.22, 0.28)

box(s2, RX, TOP + 1.26, RW, 0.5,
    "TIMETABLE WINDOWS  +  CORRIDOR CONSTRAINTS",
    fill=GREYL, line=STEEL, font=NAVY, size=9, bold=True, lw=1.0)

arrow(s2, RX + RW / 2 - 0.13, TOP + 1.82, 0.26, 0.3, fill=NAVY)

box(s2, RX, TOP + 2.18, RW, 0.82, "CP-SAT CONSTRAINT SOLVER",
    fill=NAVY, line=None, font=WHITE, size=13.5, bold=True,
    sub="Google OR-Tools  ·  proves optimality  ·  no training data",
    sub_size=8.5, sub_color="C8D4E8")

lc = RX + RW * 0.25
rc = RX + RW * 0.75
arrow(s2, lc - 0.11, TOP + 3.06, 0.22, 0.28, fill=GREEN)
arrow(s2, rc - 0.11, TOP + 3.06, 0.22, 0.28, fill=RED)

ow = RW / 2 - 0.14
box(s2, RX, TOP + 3.42, ow, 0.86, "OPTIMAL SCHEDULE",
    fill=GREENL, line=GREEN, font=GREEN, size=10, bold=True,
    sub="Conflict-free Gantt for all\nthree departments", sub_size=8,
    sub_color=GREY, lw=1.5)
box(s2, RX + ow + 0.28, TOP + 3.42, ow, 0.86, "MINIMAL CONFLICT SET",
    fill=REDL, line=RED, font=RED, size=10, bold=True,
    sub="“Requests 4, 7 and 11 cannot\ncoexist — drop one to proceed”",
    sub_size=8, sub_color=GREY, lw=1.5)

label(s2, RX, TOP + 4.42, RW, 0.24,
      "Either answer is a proof — never a guess.",
      size=9, italic=True, color=STEEL, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
#  SLIDE 3 — Technical Approach
# ══════════════════════════════════════════════════════════════
s3 = prs.slides[2]
set_text(find(s3, "Title"), "TECHNICAL APPROACH")
set_text(find(s3, "Oval"), TEAM_NAME, size=12)

body(find(s3, "TextBox 8"), [
    ("head", "Technology"),
    ("b", "Solver — Google OR-Tools CP-SAT (open source)"),
    ("b", "Backend — Python 3, FastAPI REST service"),
    ("b", "Frontend — React.js with Plotly interactive Gantt"),
    ("b", "Data — parameterised corridor scenario generator"),
    ("bg", "Runs fully offline on one laptop — no GPU, no cloud"),
    ("head", "Constraints modelled"),
    ("b", "Section exclusivity — one occupant per track section"),
    ("b", "Timetable compliance — blocks fit approved windows only"),
    ("b", "Duration bounds and inter-block corridor sequencing"),
    ("b", "Crew availability per department per day"),
    ("b", "Urgency-weighted objective — tunable at run time"),
], LX, TOP, LW, 5.1)

# ── Diagram: 5-step vertical pipeline ──
steps = [
    ("1", "INPUT", "Maintenance requests, timetable windows, corridor map", STEEL),
    ("2", "MODEL", "Encode sections, windows and priorities as CP-SAT constraints", STEEL),
    ("3", "SOLVE", "Prove the optimal assignment — or return the conflict set", NAVY),
    ("4", "VISUALISE", "Colour-coded Gantt, live priority slider, conflict panel", STEEL),
    ("5", "EXPORT", "Weekly and monthly block orders for all three departments", STEEL),
]
sh, sg = 0.78, 0.26
for i, (n, ttl, desc, col) in enumerate(steps):
    y = TOP + 0.18 + i * (sh + sg)
    hero = (col == NAVY)
    box(s3, RX + 0.55, y, RW - 0.55, sh, "",
        fill=NAVY if hero else WHITE, line=NAVY if hero else STEEL,
        font=WHITE if hero else NAVY, lw=1.5 if hero else 1.0)
    label(s3, RX + 0.72, y + 0.11, 1.5, 0.24, ttl, size=10.5, bold=True,
          color=WHITE if hero else NAVY)
    label(s3, RX + 0.72, y + 0.37, RW - 0.95, 0.34, desc, size=8.5,
          color="C8D4E8" if hero else GREY)
    box(s3, RX, y + sh / 2 - 0.19, 0.38, 0.38, n,
        fill=NAVY if hero else STEEL, line=None, font=WHITE, size=13,
        bold=True, shape=MSO_SHAPE.OVAL)
    if i < len(steps) - 1:
        arrow(s3, RX + 0.55 + (RW - 0.55) / 2 - 0.09,
              y + sh + 0.03, 0.18, 0.2)


# ══════════════════════════════════════════════════════════════
#  SLIDE 4 — Feasibility and Viability
# ══════════════════════════════════════════════════════════════
s4 = prs.slides[3]
set_text(find(s4, "Title"), "FEASIBILITY AND VIABILITY")
set_text(find(s4, "Oval"), TEAM_NAME, size=12)

body(find(s4, "TextBox 8"), [
    ("head", "Why it is buildable"),
    ("b", "OR-Tools installs with one pip command — no GPU, no licence fee"),
    ("b", "The core solver is roughly 150 lines of Python"),
    ("b", "Standard Python + React stack — nothing new to learn mid-build"),
    ("bg", "The statement lists its own inputs, so there is no data barrier"),
    ("head", "How we de-risk it"),
    ("b", "Week 1 gate: a 50-line CP-SAT model must solve a toy corridor"),
    ("b", "Demo one small division and one full corridor to prove scaling"),
    ("b", "Partial schedules still render — no all-or-nothing failure"),
    ("b", "Six members on parallel tracks: solver, UI, scenarios, domain, pitch"),
], LX, TOP, LW, 5.1)

# ── Diagram: risk -> mitigation ──
pw = RW * 0.44                       # risk column width
aw = 0.46                            # arrow lane
mx = RX + pw + aw + 0.18             # mitigation column x
mw = RX + RW - mx

hy = TOP + 0.16
label(s4, RX, hy, pw, 0.24, "RISK", size=9, bold=True,
      color=RED, align=PP_ALIGN.CENTER)
label(s4, mx, hy, mw, 0.24, "MITIGATION", size=9, bold=True,
      color=GREEN, align=PP_ALIGN.CENTER)

pairs = [
    ("No access to live TMS / BDMS / SMMS systems",
     "Adapter layer documented; parameterised scenarios are accepted method"),
    ("Railway block-working rules are unfamiliar territory",
     "One member owns the rulebook from week 1 and validates every constraint"),
    ("Only 36 hours to build at the finale",
     "Solver proven before the event; six members run independent workstreams"),
]
py, ph, pgap = hy + 0.34, 1.14, 0.34
for i, (risk, fix) in enumerate(pairs):
    y = py + i * (ph + pgap)
    box(s4, RX, y, pw, ph, risk, fill=REDL, line=RED, font=GREY,
        size=9.5, lw=1.25, pad=0.14)
    arrow(s4, RX + pw + 0.09, y + ph / 2 - 0.13, aw, 0.26,
          direction="right", fill=STEEL)
    box(s4, mx, y, mw, ph, fix, fill=GREENL, line=GREEN, font=GREY,
        size=9.5, lw=1.25, pad=0.14)


# ══════════════════════════════════════════════════════════════
#  SLIDE 5 — Impact and Benefits
# ══════════════════════════════════════════════════════════════
s5 = prs.slides[4]
set_text(find(s5, "Title"), "IMPACT AND BENEFITS")
set_text(find(s5, "Oval"), TEAM_NAME, size=12)

body(find(s5, "TextBox 8"), [
    ("head", "Operational"),
    ("b", "Clashes are resolved weeks ahead instead of on the morning itself"),
    ("b", "Every approved maintenance window gets filled, none wasted"),
    ("b", "Fewer train delays caused by blocks that overrun or collide"),
    ("b", "Each decision traces back to a named constraint — fully auditable"),
    ("head", "Reach and sustainability"),
    ("b", "Scales from one division to all 68 with no change to the code"),
    ("b", "Feeds the existing eBlock workflow — no retraining for field staff"),
    ("bg", "Zero licensing cost — the entire stack is open source"),
    ("b", "Better-planned maintenance means healthier track and safer running"),
], LX, TOP, LW, 5.1)

# ── Diagram: before / after Gantt ──
lanes = [("ENG", C_ENG), ("SIG", C_SIG), ("TRD", C_TRD)]
lane_h, lane_gap = 0.34, 0.09
lane_pitch = lane_h + lane_gap
block_h = 3 * lane_h + 2 * lane_gap        # height of one 3-lane chart
track_x, track_w = RX + 0.62, RW - 0.62


def gantt(y0, title, title_col, bars, clashes=(), foot="", foot_col=GREY):
    """Draw a 3-lane Gantt at y0. bars = [(lane_index, x_off, width)]."""
    label(s5, RX, y0, RW, 0.24, title, size=9.5, bold=True, color=title_col)
    top = y0 + 0.28
    for i, (nm, _) in enumerate(lanes):
        y = top + i * lane_pitch
        box(s5, RX, y, 0.52, lane_h, nm, fill=GREYL, line=None, font=GREY,
            size=8, bold=True, shape=MSO_SHAPE.RECTANGLE)
        box(s5, track_x, y, track_w, lane_h, "", fill="F7F9FB", line=None,
            shape=MSO_SHAPE.RECTANGLE)
    for li, xo, w in bars:
        bar(s5, track_x + xo, top + li * lane_pitch + 0.05, w,
            lane_h - 0.10, lanes[li][1])
    for xo, w in clashes:
        box(s5, track_x + xo, top - 0.03, w, block_h + 0.06, "",
            fill=None, line=RED, lw=1.5, shape=MSO_SHAPE.RECTANGLE)
    if foot:
        label(s5, RX, top + block_h + 0.07, RW, 0.22, foot,
              size=8.5, bold=True, color=foot_col)
    return top + block_h + 0.07 + 0.22


# TODAY — ENG bar1 (0.15–1.70) overlaps SIG (1.20–2.60);
#         ENG bar2 (3.05–4.35) overlaps TRD (3.45–4.95)
y = gantt(
    TOP + 0.02, "TODAY  —  each department plans on its own", RED,
    bars=[(0, 0.15, 1.55), (0, 3.05, 1.30),
          (1, 1.20, 1.40), (2, 3.45, 1.50)],
    clashes=[(1.20, 0.50), (3.45, 0.90)],
    foot="✖   blocks collide — one gets cancelled on the day",
    foot_col=RED)

# WITH RAILBLOCK — nothing overlaps
y = gantt(
    y + 0.30, "WITH RAILBLOCK  —  one optimised plan", GREEN,
    bars=[(0, 0.10, 1.45), (0, 4.05, 1.30),
          (1, 1.62, 1.25), (2, 2.94, 1.05)],
    foot="✓   zero conflicts  ·  every window used  ·  proven optimal",
    foot_col=GREEN)

# ── Stat tiles ──
stats = [("13 M+", "passengers a day"), ("68", "divisions reachable"),
         ("3", "departments unified"), ("< 2 s", "to re-solve")]
tw = (RW - 3 * 0.16) / 4
ty = y + 0.30
for i, (n, l) in enumerate(stats):
    box(s5, RX + i * (tw + 0.16), ty, tw, 0.72, n,
        fill=NAVY, line=None, font=WHITE, size=15, bold=True,
        sub=l, sub_size=7.5, sub_color="C8D4E8")


# ══════════════════════════════════════════════════════════════
#  SLIDE 6 — Research and References
# ══════════════════════════════════════════════════════════════
s6 = prs.slides[5]
set_text(find(s6, "Title"), "RESEARCH AND REFERENCES")
set_text(find(s6, "Oval"), TEAM_NAME, size=12)

body(find(s6, "TextBox 8"), [
    ("head", "Domain"),
    ("b", "Indian Railways — Block Working Rules and departmental circulars"),
    ("b", "Ministry of Railways — eBlock system documentation"),
    ("b", "NTES — National Train Enquiry System timetable data"),
    ("head", "Technical"),
    ("b", "Google OR-Tools CP-SAT — developers.google.com/optimization"),
    ("b", "Caprara, Fischetti & Toth, “Modeling and Solving the Train "
          "Timetabling Problem”, Operations Research 50(5), 2002"),
    ("b", "Perron & Furnon, OR-Tools, Google LLC, 2023"),
    ("head", "Precedent"),
    ("bg", "SIH 2025 — AI-Driven Train Induction Planning for KMRL: the same "
           "class of railway scheduling problem, solved with optimisation"),
], LX, TOP, LW, 5.1)

# ── Diagram: hub ──
label(s6, RX, TOP + 0.42, RW, 0.24, "SYSTEM AT A GLANCE", size=9,
      bold=True, color=STEEL, align=PP_ALIGN.CENTER)

cx, cy = RX + RW / 2, TOP + 2.22
box(s6, cx - 0.95, cy - 0.48, 1.9, 0.96, "CP-SAT", fill=NAVY, line=None,
    font=WHITE, size=13, bold=True, sub="constraint solver", sub_size=8,
    sub_color="C8D4E8", shape=MSO_SHAPE.OVAL)

sat = [(-2.35, -1.35, "DEMANDS", "3 departments", C_ENG),
       (0.45,  -1.35, "TIMETABLE", "approved windows", C_SIG),
       (-2.35,  0.55, "GANTT", "optimal plan", GREEN),
       (0.45,   0.55, "CONFLICTS", "minimal set", RED)]
for dx, dy, ttl, sb, col in sat:
    box(s6, cx + dx, cy + dy, 1.9, 0.68, ttl, fill=WHITE, line=col,
        font=col, size=9.5, bold=True, sub=sb, sub_size=7.5,
        sub_color=STEEL, lw=1.5)

label(s6, RX, cy + 1.5, RW, 0.5,
      "Four inputs, one solver, two possible answers — a proven schedule "
      "or the exact reason one is impossible.",
      size=9, italic=True, color=STEEL, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════
#  Remove the instructions slide (slide 7)
# ══════════════════════════════════════════════════════════════
sldIdLst = prs.element.find(qn("p:sldIdLst"))
if len(sldIdLst) > 6:
    sldIdLst.remove(sldIdLst[6])

prs.save(DEST)
print("Saved:", DEST)
print("Next: open in PowerPoint, File > Export > Create PDF, upload to portal")
